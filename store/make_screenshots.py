"""Capture Microsoft Store screenshots from the real application.

Nothing here is mocked: it builds a set of genuine sample documents, drives
the actual PySide6 window, runs a real MarkItDown conversion and grabs the
window at each stage. Store listings must show the app as it really behaves,
so the output is a true picture of it.

    python store/make_screenshots.py [--out DIR] [--size 1366x768]

Partner Center accepts 1366x768 or 1920x1080 PNG screenshots for desktop
apps; 1366x768 is the default here.

The app applies its own Qt stylesheet with the Fusion style, so it renders
almost identically on every platform — but re-running this on Windows gives
the most faithful shots for a Windows Store listing.
"""

from __future__ import annotations

import argparse
import os
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

# Render off-screen at 1x so the captures come out at exactly the asked-for
# pixel size on HiDPI machines too.
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("QT_SCALE_FACTOR", "1")
os.environ.setdefault("QT_ENABLE_HIGHDPI_SCALING", "0")

DEFAULT_OUT = ROOT / "store" / "assets" / "screenshots"


# --------------------------------------------------------------------------- #
# Sample documents — real files in real formats, so the conversion is real.
# --------------------------------------------------------------------------- #
def _write_pdf(path: Path, title: str, lines: list[str]) -> None:
    """Build a genuine PDF using Qt's own PDF writer (no extra dependency)."""
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QFont, QPageSize, QPainter, QPdfWriter

    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.A4))
    writer.setPageMargins(QMarginsF(20, 20, 20, 20))
    writer.setTitle(title)

    painter = QPainter(writer)
    heading = QFont("Helvetica", 18)
    heading.setBold(True)
    painter.setFont(heading)
    painter.drawText(0, 400, title)

    painter.setFont(QFont("Helvetica", 11))
    y = 1200
    for line in lines:
        painter.drawText(0, y, line)
        y += 500
    painter.end()


def _write_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Q3"
    ws.append(["Region", "Units", "Revenue (USD)"])
    for row in [
        ["North America", 1240, 186000],
        ["Europe", 980, 147000],
        ["Asia Pacific", 1510, 211400],
        ["Latin America", 430, 64500],
    ]:
        ws.append(row)
    wb.save(path)


def _write_pptx(path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Product Roadmap"
    slide.placeholders[1].text = (
        "Q1 — Batch conversion\nQ2 — Folder watching\nQ3 — Plugin gallery"
    )
    prs.save(path)


def build_samples(folder: Path) -> list[Path]:
    """Create one real file per major supported format."""
    folder.mkdir(parents=True, exist_ok=True)

    pdf = folder / "Q3 Financial Report.pdf"
    _write_pdf(
        pdf,
        "Q3 Financial Report",
        [
            "Revenue grew 18% quarter over quarter, led by Asia Pacific.",
            "Operating margin held steady at 22%.",
            "Headcount ended the quarter at 148.",
        ],
    )

    xlsx = folder / "Sales Figures.xlsx"
    _write_xlsx(xlsx)

    pptx = folder / "Product Roadmap.pptx"
    _write_pptx(pptx)

    html = folder / "Onboarding Guide.html"
    html.write_text(
        "<html><head><title>Onboarding Guide</title></head><body>"
        "<h1>Onboarding Guide</h1>"
        "<p>Welcome to the team. Start with the checklist below.</p>"
        "<ul><li>Set up your accounts</li><li>Read the handbook</li>"
        "<li>Meet your buddy</li></ul></body></html>",
        encoding="utf-8",
    )

    csv = folder / "Customer List.csv"
    csv.write_text(
        "Name,Company,Plan\n"
        "A. Rahman,Northwind,Team\n"
        "J. Okafor,Contoso,Business\n"
        "L. Fernandes,Fabrikam,Team\n",
        encoding="utf-8",
    )

    return [pdf, xlsx, pptx, html, csv]


# --------------------------------------------------------------------------- #
# Capture
# --------------------------------------------------------------------------- #
def _pump(app, ms: int = 250) -> None:
    """Let Qt lay out and paint before grabbing."""
    from PySide6.QtCore import QDeadlineTimer, QEventLoop

    deadline = QDeadlineTimer(ms)
    while not deadline.hasExpired():
        app.processEvents(QEventLoop.AllEvents, 20)


def _grab(window, path: Path, size: tuple[int, int]) -> None:
    pixmap = window.grab()
    image = pixmap.toImage()
    if (image.width(), image.height()) != size:
        from PySide6.QtCore import Qt

        image = image.scaled(
            size[0], size[1], Qt.IgnoreAspectRatio, Qt.SmoothTransformation
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(str(path), "PNG")
    print(f"  {path.name}  ({image.width()}x{image.height()})")


def capture(out_dir: Path, size: tuple[int, int]) -> int:
    from PySide6.QtWidgets import QApplication

    from anytomd.gui import MainWindow, apply_theme

    workdir = Path(tempfile.mkdtemp(prefix="anytomd-store-"))
    try:
        # QApplication has to exist before the PDF sample can be drawn — Qt's
        # PDF writer needs the font database.
        app = QApplication.instance() or QApplication([])
        apply_theme(app)

        samples = build_samples(workdir / "Documents")

        window = MainWindow()
        window.resize(*size)
        window.show()
        _pump(app, 400)

        print("Capturing:")
        # 1. The empty drag-and-drop state a new user sees first.
        _grab(window, out_dir / "01-drop-files.png", size)

        # 2. A batch queued up. Output stays on the default "alongside each
        #    source file" so the shot shows no machine-specific temp path.
        window._add_paths(samples)
        window.statusBar().showMessage(f"Added {len(samples)} file(s).")
        _pump(app, 300)
        _grab(window, out_dir / "02-batch-queued.png", size)

        # 3. The result of an actual conversion — really running MarkItDown.
        window._start_conversion()
        for _ in range(600):           # up to ~60s
            _pump(app, 100)
            if window._thread is None:
                break
        else:
            raise RuntimeError("Conversion did not finish in time.")
        _pump(app, 300)
        _grab(window, out_dir / "03-conversion-complete.png", size)

        produced = sorted(p.name for p in (workdir / "Documents").glob("*.md"))
        print(f"\nReal conversion produced {len(produced)} Markdown files: "
              f"{', '.join(produced)}")
        window.close()
        return 0 if produced else 1
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument(
        "--size", default="1366x768",
        help="Screenshot size accepted by Partner Center (1366x768 or 1920x1080).",
    )
    args = parser.parse_args(argv)

    try:
        width, height = (int(v) for v in args.size.lower().split("x"))
    except ValueError:
        raise SystemExit(f"--size must look like 1366x768, got {args.size!r}")

    return capture(args.out, (width, height))


if __name__ == "__main__":
    sys.exit(main())
